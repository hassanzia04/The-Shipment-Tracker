import io
import uuid
from fastapi import HTTPException, UploadFile
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
from openpyxl import load_workbook, Workbook

from sqlalchemy import or_

from app.auth.models import User
from app.masters.models import Truck, OutsourcedTruck, ProductType, RopInspectionType, BayanType, Consignee, OffloadingPoint, LoadingPort, ShippingLine
from app.masters.schemas import ExcelImportResult
from app.tenancy import company_scope


async def deactivate_truck(db: AsyncSession, truck_id: uuid.UUID) -> Truck:
    result = await db.execute(select(Truck).where(Truck.id == truck_id))
    truck = result.scalar_one_or_none()
    if not truck:
        raise HTTPException(status_code=404, detail="Truck not found")

    from app.shipments.models import Container
    from app.enums import ContainerStatus
    in_use = await db.execute(
        select(Container.id).where(
            Container.truck_id == truck_id,
            Container.status.notin_([ContainerStatus.RETURNED, ContainerStatus.CLOSED]),
        ).limit(1)
    )
    if in_use.scalar_one_or_none():
        raise HTTPException(
            status_code=409,
            detail="Cannot deactivate: truck is assigned to one or more active containers.",
        )

    truck.is_active = False
    await db.commit()
    await db.refresh(truck)
    return truck


async def delete_truck(db: AsyncSession, truck_id: uuid.UUID) -> None:
    result = await db.execute(select(Truck).where(Truck.id == truck_id))
    truck = result.scalar_one_or_none()
    if not truck:
        raise HTTPException(status_code=404, detail="Truck not found")

    from app.shipments.models import Container
    in_use = await db.execute(
        select(Container.id).where(Container.truck_id == truck_id).limit(1)
    )
    if in_use.scalar_one_or_none():
        raise HTTPException(
            status_code=409,
            detail="Cannot delete: truck is assigned to one or more containers.",
        )

    await db.delete(truck)
    await db.commit()


async def delete_outsourced_truck(db: AsyncSession, truck_id: uuid.UUID) -> None:
    result = await db.execute(select(OutsourcedTruck).where(OutsourcedTruck.id == truck_id))
    truck = result.scalar_one_or_none()
    if not truck:
        raise HTTPException(status_code=404, detail="Outsourced truck not found")

    from app.shipments.models import Container
    in_use = await db.execute(
        select(Container.id).where(Container.outsourced_truck_id == truck_id).limit(1)
    )
    if in_use.scalar_one_or_none():
        raise HTTPException(
            status_code=409,
            detail="Cannot delete: truck is assigned to one or more containers.",
        )

    await db.delete(truck)
    await db.commit()


async def _delete_simple(db: AsyncSession, model, record_id: uuid.UUID, fk_column):
    result = await db.execute(select(model).where(model.id == record_id))
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=404, detail="Record not found")

    if fk_column is not None:
        from app.shipments.models import Shipment
        in_use = await db.execute(select(Shipment.id).where(fk_column == record_id).limit(1))
        if in_use.scalar_one_or_none():
            raise HTTPException(
                status_code=409,
                detail="Cannot delete: this master is referenced by one or more shipments.",
            )

    await db.delete(record)
    await db.commit()


async def _deactivate_simple(db: AsyncSession, model, record_id: uuid.UUID, fk_column, fk_model):
    result = await db.execute(select(model).where(model.id == record_id))
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=404, detail="Record not found")

    if fk_column is not None and fk_model is not None:
        from app.enums import ShipmentStage
        in_use = await db.execute(
            select(fk_model.id).where(
                fk_column == record_id,
                fk_model.current_stage.notin_([ShipmentStage.COMPLETED]),
            ).limit(1)
        )
        if in_use.scalar_one_or_none():
            raise HTTPException(
                status_code=409,
                detail="Cannot deactivate: record is referenced by one or more active shipments.",
            )

    record.is_active = False
    await db.commit()
    await db.refresh(record)
    return record


async def list_trucks(db: AsyncSession) -> list[Truck]:
    result = await db.execute(select(Truck).where(Truck.is_active == True).order_by(Truck.plate_number))
    return list(result.scalars().all())


async def create_truck(db: AsyncSession, plate: str, driver: str, contractor: str, nationality: str) -> Truck:
    truck = Truck(plate_number=plate, driver_name=driver, contractor=contractor, nationality=nationality)
    db.add(truck)
    await db.commit()
    await db.refresh(truck)
    return truck


async def import_trucks_excel(db: AsyncSession, file: UploadFile) -> ExcelImportResult:
    content = await file.read()
    wb = load_workbook(io.BytesIO(content), read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    inserted, skipped, errors = 0, 0, []
    for i, row in enumerate(rows, start=2):
        if not any(row):
            continue
        try:
            plate, driver, contractor, nationality = (str(c).strip() if c else "" for c in row[:4])
            if not plate:
                errors.append(f"Row {i}: plate number is required")
                continue

            existing = await db.execute(select(Truck).where(Truck.plate_number == plate))
            if existing.scalar_one_or_none():
                skipped += 1
                continue

            db.add(Truck(plate_number=plate, driver_name=driver, contractor=contractor, nationality=nationality))
            inserted += 1
        except Exception as e:
            errors.append(f"Row {i}: {e}")

    await db.commit()
    return ExcelImportResult(inserted=inserted, skipped=skipped, errors=errors)


def build_trucks_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Trucks"
    ws.append(["Plate Number", "Driver Name", "Contractor", "Nationality"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _list_simple(db: AsyncSession, model, actor: User | None = None):
    q = select(model).where(model.is_active == True).order_by(model.name)
    # Company-scoped masters: customer users see shared rows (NULL) + their own
    if actor is not None and hasattr(model, "company_id"):
        scope = company_scope(actor)
        if scope is not None:
            q = q.where(or_(model.company_id == None, model.company_id == scope))
    result = await db.execute(q)
    return list(result.scalars().all())


def _creation_company(model, actor: User | None) -> "uuid.UUID | None":
    """Rows created by customer users belong to their company; rows created by
    internal users are shared (visible to everyone)."""
    if actor is not None and hasattr(model, "company_id"):
        return company_scope(actor)
    return None


async def _assert_name_available(db: AsyncSession, model, name: str, company_id) -> bool:
    """A name clashes with shared rows and rows of the same company, but is free
    to repeat across different companies. Returns True if taken."""
    q = select(model.id).where(model.name == name)
    if hasattr(model, "company_id"):
        if company_id is None:
            q = q.where(model.company_id == None)
        else:
            q = q.where(or_(model.company_id == None, model.company_id == company_id))
    return (await db.execute(q)).scalar_one_or_none() is not None


async def _create_simple(db: AsyncSession, model, name: str, actor: User | None = None):
    creation_company = _creation_company(model, actor)
    if await _assert_name_available(db, model, name, creation_company):
        raise HTTPException(status_code=400, detail=f"{name} already exists")
    obj = model(name=name)
    if hasattr(model, "company_id"):
        obj.company_id = creation_company
    db.add(obj)
    await db.commit()
    await db.refresh(obj)
    return obj


async def _import_simple_excel(db: AsyncSession, model, file: UploadFile, actor: User | None = None) -> ExcelImportResult:
    content = await file.read()
    wb = load_workbook(io.BytesIO(content), read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    creation_company = _creation_company(model, actor)

    inserted, skipped, errors = 0, 0, []
    for i, row in enumerate(rows, start=2):
        if not row or not row[0]:
            continue
        try:
            name = str(row[0]).strip()
            if await _assert_name_available(db, model, name, creation_company):
                skipped += 1
                continue
            obj = model(name=name)
            if hasattr(model, "company_id"):
                obj.company_id = creation_company
            db.add(obj)
            inserted += 1
        except Exception as e:
            errors.append(f"Row {i}: {e}")

    await db.commit()
    return ExcelImportResult(inserted=inserted, skipped=skipped, errors=errors)


def build_simple_template(sheet_name: str, column: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append([column])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# Convenience wrappers for each simple master
# (company-scoped masters take the actor; global masters ignore it)
async def list_product_types(db, actor=None): return await _list_simple(db, ProductType, actor)
async def create_product_type(db, name, actor=None): return await _create_simple(db, ProductType, name, actor)
async def import_product_types(db, file, actor=None): return await _import_simple_excel(db, ProductType, file, actor)

async def list_rop_types(db): return await _list_simple(db, RopInspectionType)
async def create_rop_type(db, name): return await _create_simple(db, RopInspectionType, name)
async def import_rop_types(db, file): return await _import_simple_excel(db, RopInspectionType, file)

async def list_offloading_points(db, actor=None): return await _list_simple(db, OffloadingPoint, actor)
async def create_offloading_point(db, name, actor=None): return await _create_simple(db, OffloadingPoint, name, actor)
async def import_offloading_points(db, file, actor=None): return await _import_simple_excel(db, OffloadingPoint, file, actor)

async def list_loading_ports(db): return await _list_simple(db, LoadingPort)
async def create_loading_port(db, name): return await _create_simple(db, LoadingPort, name)
async def import_loading_ports(db, file): return await _import_simple_excel(db, LoadingPort, file)

async def list_shipping_lines(db): return await _list_simple(db, ShippingLine)
async def create_shipping_line(db, name): return await _create_simple(db, ShippingLine, name)
async def import_shipping_lines(db, file): return await _import_simple_excel(db, ShippingLine, file)

async def list_bayan_types(db): return await _list_simple(db, BayanType)
async def create_bayan_type(db, name): return await _create_simple(db, BayanType, name)
async def import_bayan_types(db, file): return await _import_simple_excel(db, BayanType, file)

async def list_consignees(db, actor=None): return await _list_simple(db, Consignee, actor)
async def create_consignee(db, name, actor=None): return await _create_simple(db, Consignee, name, actor)
async def import_consignees(db, file, actor=None): return await _import_simple_excel(db, Consignee, file, actor)

async def list_outsourced_trucks(db: AsyncSession) -> list[OutsourcedTruck]:
    result = await db.execute(select(OutsourcedTruck).where(OutsourcedTruck.is_active == True).order_by(OutsourcedTruck.plate_number))
    return list(result.scalars().all())

async def create_outsourced_truck(db: AsyncSession, plate: str, driver: str, contractor: str, nationality: str) -> OutsourcedTruck:
    truck = OutsourcedTruck(plate_number=plate, driver_name=driver, contractor=contractor, nationality=nationality)
    db.add(truck)
    await db.commit()
    await db.refresh(truck)
    return truck

async def import_outsourced_trucks_excel(db: AsyncSession, file) -> ExcelImportResult:
    return await _import_trucks_excel_generic(db, OutsourcedTruck, file)

async def deactivate_outsourced_truck(db: AsyncSession, truck_id: uuid.UUID) -> OutsourcedTruck:
    result = await db.execute(select(OutsourcedTruck).where(OutsourcedTruck.id == truck_id))
    truck = result.scalar_one_or_none()
    if not truck:
        raise HTTPException(status_code=404, detail="Outsourced truck not found")

    from app.shipments.models import Container
    from app.enums import ContainerStatus
    in_use = await db.execute(
        select(Container.id).where(
            Container.outsourced_truck_id == truck_id,
            Container.status.notin_([ContainerStatus.RETURNED, ContainerStatus.CLOSED, ContainerStatus.OFFLOADED]),
        ).limit(1)
    )
    if in_use.scalar_one_or_none():
        raise HTTPException(
            status_code=409,
            detail="Cannot deactivate: outsourced truck is assigned to one or more active containers.",
        )

    truck.is_active = False
    await db.commit()
    await db.refresh(truck)
    return truck


async def _import_trucks_excel_generic(db: AsyncSession, model, file) -> ExcelImportResult:
    from openpyxl import load_workbook
    import io
    content = await file.read()
    wb = load_workbook(io.BytesIO(content), read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    inserted, skipped, errors = 0, 0, []
    for i, row in enumerate(rows, start=2):
        if not any(row):
            continue
        try:
            plate, driver, contractor, nationality = (str(c).strip() if c else "" for c in row[:4])
            if not plate:
                errors.append(f"Row {i}: plate number is required")
                continue
            existing = await db.execute(select(model).where(model.plate_number == plate))
            if existing.scalar_one_or_none():
                skipped += 1
                continue
            db.add(model(plate_number=plate, driver_name=driver, contractor=contractor, nationality=nationality))
            inserted += 1
        except Exception as e:
            errors.append(f"Row {i}: {e}")

    await db.commit()
    return ExcelImportResult(inserted=inserted, skipped=skipped, errors=errors)


# Deactivation wrappers — each checks for active FK references before setting is_active=False
async def deactivate_product_type(db, record_id):
    from app.shipments.models import Shipment
    return await _deactivate_simple(db, ProductType, record_id, Shipment.product_type_id, Shipment)

async def deactivate_rop_type(db, record_id):
    from app.shipments.models import Shipment
    return await _deactivate_simple(db, RopInspectionType, record_id, Shipment.rop_inspection_type_id, Shipment)

async def deactivate_offloading_point(db, record_id):
    from app.shipments.models import Shipment
    return await _deactivate_simple(db, OffloadingPoint, record_id, Shipment.offloading_point_id, Shipment)

async def deactivate_loading_port(db, record_id):
    from app.shipments.models import Shipment
    return await _deactivate_simple(db, LoadingPort, record_id, Shipment.loading_port_id, Shipment)

async def deactivate_shipping_line(db, record_id):
    from app.shipments.models import Shipment
    return await _deactivate_simple(db, ShippingLine, record_id, Shipment.shipping_line_id, Shipment)


# Delete wrappers
async def delete_product_type(db, record_id):
    from app.shipments.models import Shipment
    return await _delete_simple(db, ProductType, record_id, Shipment.product_type_id)

async def delete_offloading_point(db, record_id):
    from app.shipments.models import Shipment
    return await _delete_simple(db, OffloadingPoint, record_id, Shipment.offloading_point_id)

async def delete_loading_port(db, record_id):
    from app.shipments.models import Shipment
    return await _delete_simple(db, LoadingPort, record_id, Shipment.loading_port_id)

async def delete_shipping_line(db, record_id):
    from app.shipments.models import Shipment
    return await _delete_simple(db, ShippingLine, record_id, Shipment.shipping_line_id)

async def delete_bayan_type(db, record_id):
    from app.shipments.models import Shipment
    return await _delete_simple(db, BayanType, record_id, Shipment.bayan_type_id)

async def delete_consignee(db, record_id):
    from app.shipments.models import Shipment
    return await _delete_simple(db, Consignee, record_id, Shipment.consignee_id)

async def deactivate_bayan_type(db, record_id):
    from app.shipments.models import Shipment
    return await _deactivate_simple(db, BayanType, record_id, Shipment.bayan_type_id, Shipment)

async def deactivate_consignee(db, record_id):
    from app.shipments.models import Shipment
    return await _deactivate_simple(db, Consignee, record_id, Shipment.consignee_id, Shipment)
